"""Tests for Markdown to .docx export (service + router)."""

import io
from types import SimpleNamespace

import httpx
import pytest
from docx import Document
from docx.document import Document as DocxDocument

from backend.app import create_app
from backend.core.settings import Environment
from backend.dependencies import get_app_settings
from backend.services.docx_export import (
    DOCX_MEDIA_TYPE,
    export_filename,
    markdown_to_docx,
)

_MARKDOWN = """# Overview

The project ships in **May** [1] and uses `FastAPI`.

## Risks

- Budget *overrun* risk [2]
- Staffing gap
  - Nested item

1. First step
2. Second step

| Area | Owner |
| --- | --- |
| Backend | Team A |

> Key decision recorded.

```mermaid
flowchart LR
  A --> B
```

See [the portal](https://example.com).
"""


def _open(payload: bytes) -> DocxDocument:
    return Document(io.BytesIO(payload))


def test_markdown_to_docx_renders_blocks() -> None:
    document = _open(markdown_to_docx(_MARKDOWN, title="Project doc"))
    paragraphs = [(p.style.name if p.style else "", p.text) for p in document.paragraphs]

    assert paragraphs[0] == ("Title", "Project doc")
    assert ("Heading 1", "Overview") in paragraphs
    assert ("Heading 2", "Risks") in paragraphs
    assert ("List Bullet", "Budget overrun risk [2]") in paragraphs
    assert ("List Bullet 2", "Nested item") in paragraphs
    assert ("List Number", "Second step") in paragraphs
    assert ("Quote", "Key decision recorded.") in paragraphs
    assert ("Normal", "Diagram (Mermaid source):") in paragraphs
    assert ("Normal", "  A --> B") in paragraphs
    assert ("Normal", "See the portal (https://example.com).") in paragraphs

    body = next(p for p in document.paragraphs if p.text.startswith("The project"))
    assert body.text == "The project ships in May [1] and uses FastAPI."
    assert [r.text for r in body.runs if r.bold] == ["May"]

    table = document.tables[0]
    assert [[c.text for c in row.cells] for row in table.rows] == [
        ["Area", "Owner"],
        ["Backend", "Team A"],
    ]


def test_export_filename_sanitizes() -> None:
    assert export_filename("Q3 plan: v2/final", "docx") == "Q3-plan-v2-final.docx"
    assert export_filename("///", "md") == "document.md"


@pytest.fixture
def client() -> httpx.AsyncClient:
    app = create_app()
    app.dependency_overrides[get_app_settings] = lambda: SimpleNamespace(
        environment=Environment.LOCAL,
        observability=SimpleNamespace(log_level="INFO"),
    )
    return httpx.AsyncClient(transport=httpx.ASGITransport(app=app), base_url="http://test")


async def test_export_docx_route_returns_attachment(client: httpx.AsyncClient) -> None:
    async with client:
        response = await client.post(
            "/api/export/docx", json={"title": "Project doc", "markdown": "# Hi\n\nBody"}
        )
    assert response.status_code == 200
    assert response.headers["content-type"] == DOCX_MEDIA_TYPE
    assert 'filename="Project-doc.docx"' in response.headers["content-disposition"]
    assert [p.text for p in _open(response.content).paragraphs] == ["Project doc", "Hi", "Body"]


async def test_export_docx_route_rejects_empty_markdown(client: httpx.AsyncClient) -> None:
    async with client:
        response = await client.post("/api/export/docx", json={"markdown": ""})
    assert response.status_code == 422
