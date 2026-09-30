"""Markdown to Word (.docx) rendering for exported answers.

Covers the Markdown subset the synthesis prompts produce: ATX
headings, paragraphs, bullet and numbered lists, pipe tables, block
quotes, fenced code blocks (Mermaid diagrams are kept as their source
text), and inline bold / italic / code / links.
"""

import io
import re

from docx import Document
from docx.document import Document as DocxDocument
from docx.shared import Pt
from docx.text.paragraph import Paragraph

DOCX_MEDIA_TYPE = (
    "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
)

_HEADING = re.compile(r"^(#{1,6})\s+(.*?)\s*#*\s*$")
_BULLET = re.compile(r"^(\s*)[-*+]\s+(.*)$")
_NUMBERED = re.compile(r"^(\s*)\d+[.)]\s+(.*)$")
_FENCE = re.compile(r"^\s*(```|~~~)\s*([\w-]*)\s*$")
_RULE = re.compile(r"^\s*([-*_])(\s*\1){2,}\s*$")
_TABLE_SEPARATOR = re.compile(r"^\s*\|?\s*:?-{2,}:?\s*(\|\s*:?-{2,}:?\s*)*\|?\s*$")
_INLINE = re.compile(
    r"(\*\*(?P<bold>.+?)\*\*"
    r"|__(?P<bold2>.+?)__"
    r"|`(?P<code>[^`]+)`"
    r"|\[(?P<link_text>[^\]]+)\]\((?P<link_url>[^)\s]+)\)"
    r"|(?<![\w*])\*(?P<italic>[^*\s][^*]*?)\*(?![\w*])"
    r"|(?<!\w)_(?P<italic2>[^_\s][^_]*?)_(?!\w))"
)
_CODE_FONT = "Consolas"
_FILENAME_UNSAFE = re.compile(r"[^A-Za-z0-9._-]+")


def export_filename(title: str, extension: str) -> str:
    """Return a filesystem-safe download name for ``title``."""
    stem = _FILENAME_UNSAFE.sub("-", title).strip("-.") or "document"
    return f"{stem[:100]}.{extension}"


def _add_inline(paragraph: Paragraph, text: str) -> None:
    position = 0
    for match in _INLINE.finditer(text):
        if match.start() > position:
            paragraph.add_run(text[position : match.start()])
        bold = match.group("bold") or match.group("bold2")
        italic = match.group("italic") or match.group("italic2")
        code = match.group("code")
        link_text = match.group("link_text")
        if bold is not None:
            paragraph.add_run(bold).bold = True
        elif italic is not None:
            paragraph.add_run(italic).italic = True
        elif code is not None:
            paragraph.add_run(code).font.name = _CODE_FONT
        elif link_text is not None:
            paragraph.add_run(f"{link_text} ({match.group('link_url')})")
        position = match.end()
    if position < len(text):
        paragraph.add_run(text[position:])


def _split_row(line: str) -> list[str]:
    stripped = line.strip()
    if stripped.startswith("|"):
        stripped = stripped[1:]
    if stripped.endswith("|"):
        stripped = stripped[:-1]
    return [cell.strip() for cell in stripped.split("|")]


def _add_table(document: DocxDocument, rows: list[list[str]]) -> None:
    width = max(len(row) for row in rows)
    table = document.add_table(rows=len(rows), cols=width)
    table.style = "Table Grid"
    for row_index, row in enumerate(rows):
        for col_index in range(width):
            cell = table.cell(row_index, col_index)
            value = row[col_index] if col_index < len(row) else ""
            paragraph = cell.paragraphs[0]
            _add_inline(paragraph, value)
            if row_index == 0:
                for run in paragraph.runs:
                    run.bold = True


def _add_code(document: DocxDocument, lines: list[str], language: str) -> None:
    if language == "mermaid":
        document.add_paragraph("Diagram (Mermaid source):").runs[0].italic = True
    for line in lines or [""]:
        run = document.add_paragraph().add_run(line)
        run.font.name = _CODE_FONT
        run.font.size = Pt(9)


def _list_style(indent: str, base: str) -> str:
    return base if len(indent.expandtabs(4)) < 2 else f"{base} 2"


def markdown_to_docx(markdown: str, *, title: str) -> bytes:
    """Render ``markdown`` as a .docx document titled ``title``."""
    document = Document()
    document.add_heading(title, level=0)
    lines = markdown.replace("\r\n", "\n").split("\n")
    paragraph_buffer: list[str] = []

    def flush_paragraph() -> None:
        if paragraph_buffer:
            _add_inline(document.add_paragraph(), " ".join(paragraph_buffer))
            paragraph_buffer.clear()

    index = 0
    while index < len(lines):
        line = lines[index]
        fence = _FENCE.match(line)
        if fence is not None:
            flush_paragraph()
            marker, language = fence.group(1), fence.group(2).lower()
            code_lines: list[str] = []
            index += 1
            while index < len(lines) and not lines[index].strip().startswith(marker):
                code_lines.append(lines[index])
                index += 1
            _add_code(document, code_lines, language)
            index += 1
            continue
        if (
            line.lstrip().startswith("|")
            and index + 1 < len(lines)
            and _TABLE_SEPARATOR.match(lines[index + 1]) is not None
        ):
            flush_paragraph()
            rows = [_split_row(line)]
            index += 2
            while index < len(lines) and lines[index].lstrip().startswith("|"):
                rows.append(_split_row(lines[index]))
                index += 1
            _add_table(document, rows)
            continue
        index += 1
        if not line.strip():
            flush_paragraph()
            continue
        heading = _HEADING.match(line)
        if heading is not None:
            flush_paragraph()
            _add_inline(
                document.add_heading(level=min(len(heading.group(1)), 9)),
                heading.group(2),
            )
            continue
        if _RULE.match(line) is not None:
            flush_paragraph()
            continue
        bullet = _BULLET.match(line)
        if bullet is not None:
            flush_paragraph()
            style = _list_style(bullet.group(1), "List Bullet")
            _add_inline(document.add_paragraph(style=style), bullet.group(2))
            continue
        numbered = _NUMBERED.match(line)
        if numbered is not None:
            flush_paragraph()
            style = _list_style(numbered.group(1), "List Number")
            _add_inline(document.add_paragraph(style=style), numbered.group(2))
            continue
        if line.lstrip().startswith(">"):
            flush_paragraph()
            quote = line.lstrip()[1:].strip()
            _add_inline(document.add_paragraph(style="Quote"), quote)
            continue
        paragraph_buffer.append(line.strip())
    flush_paragraph()

    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()
